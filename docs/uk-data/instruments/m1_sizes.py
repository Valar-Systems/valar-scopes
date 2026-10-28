"""M1 stage B (docs/uk-data.md A2.4): BMPK v1 package sizes for UK centres under §2.2's UK rules.

A measurement, not the builder. The geometry steps are written from docs/basemap-package-format.md (§2, §4-§7, §13);
the bytes come from the format owner's encoder (tools/weatherscope/basemap/bmpk.py) and every package is checked by
its reference validator (validate.py). No coordinate is printed.

Usage: python -I m1_sizes.py <m1src> <m1work> <repo> <airports.csv> <out.json> [--controls] [--eps0] [--nodrop] [--nochain]
"""
import csv
import glob
import hashlib
import json
import math
import os
import pickle
import subprocess
import sys
import time
import urllib.request

import numpy as np
import openpyxl
import shapefile
from PIL import Image
from pyproj import CRS, Transformer
from shapely import STRtree, points as shp_points
from shapely.geometry import Point, shape as geo_shape, Polygon

args = [a for a in sys.argv[1:] if not a.startswith("--")]
FLAGS = {a for a in sys.argv[1:] if a.startswith("--")}
SRC, WORK, REPO, APTS, OUT = args
sys.path.insert(0, os.path.join(REPO, "tools", "weatherscope", "basemap"))
import bmpk  # noqa: E402
import validate as bmpk_validate  # noqa: E402

K = 111.0
CAP = 32 * 1024
F_LIMIT = 20 * 1024
UA = {"User-Agent": "BlipscopeBasemapM1/1 (+https://github.com/Valar-Systems/valar-scopes)"}
GB_CENTRES = ["LHR", "MAN", "BHX", "NCL", "EDI", "GLA", "ABZ", "INV", "CWL", "NQY"]
NI_CENTRES = ["BFS", "LDY"]
RADII = [(161, 671.0), (222, 925.0), (97, 100.584)]   # (km, eps m): graded 161/222, printed 97
VARIANTS = ["i", "ii", "iii", "iv"]


def q2(v):
    return math.floor(v * 100 + 0.5) / 100


# ---------------------------------------------------------------------------------------------- the sources, once
def load_gb():
    p = os.path.join(WORK, "geo_gb.pkl")
    if os.path.exists(p):
        return pickle.load(open(p, "rb"))
    tr = Transformer.from_crs("EPSG:27700", "EPSG:4326", always_xy=True)

    def ll(pts):
        a = np.asarray(pts, dtype=float)
        lon, lat = tr.transform(a[:, 0], a[:, 1])
        return np.column_stack([lon, lat])

    out = {}
    for layer in ("roads", "rivers", "water"):
        feats = pickle.load(open(os.path.join(WORK, layer + ".pkl"), "rb"))
        for f in feats:
            f["ll"] = ll(f["pts"])
            del f["pts"]
        out[layer] = feats
    for layer in ("places", "hills"):
        feats = pickle.load(open(os.path.join(WORK, layer + ".pkl"), "rb"))
        xs = np.array([f["x"] for f in feats]); ys = np.array([f["y"] for f in feats])
        lon, lat = tr.transform(xs, ys)
        for f, a, b in zip(feats, lon, lat):
            f["lon"], f["lat"] = float(a), float(b)
        out[layer] = feats
    # populations: E&W BUA polygons with ONS counts (London absent from the table -> no count), Scotland NRS Popcount
    ons = {}
    wb = openpyxl.load_workbook(os.path.join(SRC, "ons", "ons-bua-census2021.xlsx"), read_only=True)
    for sheet in ("1c", "1d"):
        rows = list(wb[sheet].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and "BUA code" in [str(c).strip() for c in r if c is not None])
        hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
        ci, cc = hdr.index("BUA code"), next(i for i, h in enumerate(hdr) if h.startswith("Count"))
        for r in rows[hi + 1:]:
            if r and r[ci] and isinstance(r[cc], (int, float)):
                ons[str(r[ci]).strip()] = int(r[cc])
    polys, pops = [], []
    for gp in sorted(glob.glob(os.path.join(SRC, "ons-bua", "*.geojson"))):
        for ft in json.load(open(gp, encoding="utf-8"))["features"]:
            polys.append(geo_shape(ft["geometry"]))
            pops.append(ons.get(ft["properties"]["BUA22CD"], 0))
    nrs_p = glob.glob(os.path.join(SRC, "nrs", "**", "*.shp"), recursive=True)[0]
    r = shapefile.Reader(nrs_p, encoding="utf-8")
    nrs_polys, nrs_pops = [], []
    for sr in r.iterShapeRecords():
        nrs_polys.append(geo_shape(sr.shape.__geo_interface__))
        nrs_pops.append(int(sr.record["Popcount"] or 0))
    tree_b, tree_n = STRtree(polys), STRtree(nrs_polys)
    stats = {"ons_counts": len(ons), "bua_polygons": len(polys), "nrs_polygons": len(nrs_polys), "pop_from": {}}
    for f in out["places"]:
        pt = Point(f["x"], f["y"])
        pop, src = 0, "none"
        hit = [i for i in tree_n.query(pt) if nrs_polys[i].covers(pt)]
        if hit:
            pop, src = nrs_pops[min(hit)], "nrs"
        else:
            hit = [i for i in tree_b.query(pt) if polys[i].covers(pt)]
            if hit and pops[min(hit)]:
                pop, src = pops[min(hit)], "ons"
            elif hit:
                src = "bua-without-count"
        f["pop"] = pop
        stats["pop_from"][src] = stats["pop_from"].get(src, 0) + 1
    out["pop_stats"] = stats
    pickle.dump(out, open(p, "wb"), protocol=4)
    return out


def load_ni():
    p = os.path.join(WORK, "geo_ni.pkl")
    if os.path.exists(p):
        return pickle.load(open(p, "rb"))

    def reader(path):
        cpg = path[:-4] + ".cpg"
        enc = open(cpg).read().strip() if os.path.exists(cpg) else None
        ldid = open(path[:-4] + ".dbf", "rb").read(32)[29]
        if enc is None:
            if ldid not in (0x03, 0x57):
                raise SystemExit(f"{path}: unexpected DBF language driver {ldid:#x}")
            enc = "cp1252"
        return shapefile.Reader(path, encoding=enc), CRS.from_wkt(open(path[:-4] + ".prj").read())

    out = {"roads": [], "water": [], "places": []}
    rp = glob.glob(os.path.join(SRC, "osni", "*Transport_lines.shp"))[0]
    r, crs = reader(rp)
    tr = Transformer.from_crs(crs, "EPSG:4326", always_xy=True)
    for ri, sr in enumerate(r.iterShapeRecords()):
        cls = {"MOTORWAY": 1, "A_CLASS": 2}.get(sr.record["TEMA"])
        if not cls:
            continue
        pts = sr.shape.points
        parts = list(sr.shape.parts) + [len(pts)]
        for pi in range(len(parts) - 1):
            a = np.asarray(pts[parts[pi]:parts[pi + 1]], dtype=float)
            lon, lat = tr.transform(a[:, 0], a[:, 1])
            out["roads"].append({"key": ("osni/" + os.path.basename(rp), ri, pi), "cls": cls, "primary": True, "num": "",
                                 "a": None, "b": None, "ll": np.column_stack([lon, lat]), "ni": True})
    lp = glob.glob(os.path.join(SRC, "daera", "*.shp"))[0]
    r, crs = reader(lp)
    tr = Transformer.from_crs(crs, "EPSG:4326", always_xy=True)
    for ri, sr in enumerate(r.iterShapeRecords()):
        if float(sr.record["squ_metre"] or 0) < 1_500_000:
            continue
        pts = sr.shape.points
        parts = list(sr.shape.parts) + [len(pts)]
        for pi in range(len(parts) - 1):
            ring = pts[parts[pi]:parts[pi + 1]]
            if len(ring) > 1 and tuple(ring[0]) == tuple(ring[-1]):
                ring = ring[:-1]
            s = sum(x0 * y1 - x1 * y0 for (x0, y0), (x1, y1) in zip(ring, ring[1:] + ring[:1]))
            if s < 0 and len(ring) >= 3:   # outer (clockwise); holes dropped (RD3)
                a = np.asarray(ring, dtype=float)
                lon, lat = tr.transform(a[:, 0], a[:, 1])
                out["water"].append({"key": ("daera/" + os.path.basename(lp), ri, pi), "ll": np.column_stack([lon, lat]), "ni": True})
    # places: OSNI gazetteer, populations from the NISRA 2015 limits containing the point, joined to CT0046 by code
    sp = glob.glob(os.path.join(SRC, "nisra", "*.shp"))[0]
    r, scrs = reader(sp)
    lim, codes = [], []
    for sr in r.iterShapeRecords():
        lim.append(geo_shape(sr.shape.__geo_interface__)); codes.append(str(sr.record["Code"]).strip())
    wb = openpyxl.load_workbook(os.path.join(SRC, "nisra", "ct0046.xlsx"), read_only=True)
    ct = {}
    rows = list(wb["CT0046"].iter_rows(values_only=True))
    hi = next(i for i, row in enumerate(rows) if row and any(str(c).strip().lower().startswith("geography code") for c in row if c))
    hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
    ci = next(i for i, h in enumerate(hdr) if h.lower().startswith("geography code"))
    ui = next(i for i, h in enumerate(hdr) if h.lower().startswith("usual residents"))
    for row in rows[hi + 1:]:
        if row and row[ci] and isinstance(row[ui], (int, float)):
            ct[str(row[ci]).strip()] = int(row[ui])
    tree = STRtree(lim)
    gaz_crs = scrs   # the gazetteer states no CRS; OSNI publishes its gazetteers on the Irish Grid, as the limits are
    tr = Transformer.from_crs(gaz_crs, "EPSG:4326", always_xy=True)
    stats = {"ct0046_codes": len(ct), "limits": len(lim), "pop_from": {}}
    with open(os.path.join(SRC, "osni", "osni-gazetteer-place-names.csv"), encoding="utf-8-sig") as f:
        for ri, row in enumerate(csv.DictReader(f)):
            x, y = float(row["X"]), float(row["Y"])
            pt = Point(x, y)
            hit = [i for i in tree.query(pt) if lim[i].covers(pt)]
            pop = ct.get(codes[min(hit)], 0) if hit else 0
            src = "ct0046" if pop else ("limit-without-count" if hit else "none")
            stats["pop_from"][src] = stats["pop_from"].get(src, 0) + 1
            lon, lat = tr.transform(x, y)
            out["places"].append({"key": ("osni/gazetteer", ri), "id": row["OBJECTID"], "name": row["PLACENAME"].strip(),
                                  "lon": float(lon), "lat": float(lat), "pop": pop, "ni": True})
    out["pop_stats"] = stats
    pickle.dump(out, open(p, "wb"), protocol=4)
    return out


# ---------------------------------------------------------------------------------------------- geometry (§2, §6, §13)
def clip_seg(x0, y0, x1, y1, R):
    """Liang-Barsky against [-R, R]^2; returns (t0, t1) or None."""
    t0, t1 = 0.0, 1.0
    dx, dy = x1 - x0, y1 - y0
    for p, q in ((-dx, x0 + R), (dx, R - x0), (-dy, y0 + R), (dy, R - y0)):
        if p == 0:
            if q < 0:
                return None
        else:
            t = q / p
            if p < 0:
                if t > t1:
                    return None
                t0 = max(t0, t)
            else:
                if t < t0:
                    return None
                t1 = min(t1, t)
    return t0, t1


def clip_line(pts, R):
    """RD5: each in-square run is its own polyline, in along-line order."""
    runs, cur = [], None
    for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
        c = clip_seg(x0, y0, x1, y1, R)
        if c is None:
            if cur:
                runs.append(cur)
            cur = None
            continue
        t0, t1 = c
        a = (x0 + t0 * (x1 - x0), y0 + t0 * (y1 - y0))
        b = (x0 + t1 * (x1 - x0), y0 + t1 * (y1 - y0))
        if cur is None or t0 > 0:
            if cur:
                runs.append(cur)
            cur = [a]
        cur.append(b)
        if t1 < 1:
            runs.append(cur)
            cur = None
    if cur:
        runs.append(cur)
    return runs


def clip_ring(pts, R):
    """Sutherland-Hodgman against the four edges; the ring stays closed (implicitly)."""
    out = list(pts)
    for inside, inter in (
        (lambda p: p[0] >= -R, lambda a, b: (-R, a[1] + (b[1] - a[1]) * (-R - a[0]) / (b[0] - a[0]))),
        (lambda p: p[0] <= R, lambda a, b: (R, a[1] + (b[1] - a[1]) * (R - a[0]) / (b[0] - a[0]))),
        (lambda p: p[1] >= -R, lambda a, b: (a[0] + (b[0] - a[0]) * (-R - a[1]) / (b[1] - a[1]), -R)),
        (lambda p: p[1] <= R, lambda a, b: (a[0] + (b[0] - a[0]) * (R - a[1]) / (b[1] - a[1]), R)),
    ):
        if not out:
            break
        src, out = out, []
        prev = src[-1]
        for cur in src:
            if inside(cur):
                if not inside(prev):
                    out.append(inter(prev, cur))
                out.append(cur)
            elif inside(prev):
                out.append(inter(prev, cur))
            prev = cur
    return out


def seg_dist(p, a, b):
    """RD7: perpendicular distance to the SEGMENT, clamped at its ends; RD6: zero-length chord -> distance to point."""
    ax, ay = a; bx, by = b; px, py = p
    dx, dy = bx - ax, by - ay
    L = dx * dx + dy * dy
    if L == 0:
        return math.hypot(px - ax, py - ay)
    t = max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / L))
    return math.hypot(px - (ax + t * dx), py - (ay + t * dy))


def dp(pts, eps):
    """Douglas-Peucker; split at the FIRST vertex at the maximum distance; keep when distance > eps (RD7)."""
    n = len(pts)
    if n <= 2:
        return list(pts)
    keep = [False] * n
    keep[0] = keep[-1] = True
    stack = [(0, n - 1)]
    while stack:
        i, j = stack.pop()
        dmax, idx = -1.0, -1
        for k in range(i + 1, j):
            d = seg_dist(pts[k], pts[i], pts[j])
            if d > dmax:
                dmax, idx = d, k
        if idx >= 0 and dmax > eps:
            keep[idx] = True
            stack.append((idx, j))
            stack.append((i, idx))
    return [p for p, k in zip(pts, keep) if k]


def quant(pts_km, R_mm):
    out = []
    for e, s in pts_km:
        qx = math.floor(e * 1e6 / R_mm * 32767 + 0.5)
        qy = math.floor(s * 1e6 / R_mm * 32767 + 0.5)
        if abs(qx) > 32767 or abs(qy) > 32767:
            raise SystemExit("a vertex outside the square after clipping: a bug, never clamped (format §2)")
        if not out or out[-1] != (qx, qy):
            out.append((qx, qy))
    return out


class Frame:
    def __init__(self, lat, lon, r_km):
        self.lat0, self.lon0 = q2(lat), q2(lon)
        self.c0 = max(math.cos(self.lat0 * (math.pi / 180)), 0.01)
        self.R = float(r_km)
        self.R_mm = int(r_km) * 1_000_000
        self.dlat, self.dlon = self.R / K, self.R / (K * self.c0)

    def es(self, ll):
        return [((lon - self.lon0) * K * self.c0, (self.lat0 - lat) * K) for lon, lat in ll]

    def bbox_hit(self, ll):
        return (ll[:, 1].max() >= self.lat0 - self.dlat and ll[:, 1].min() <= self.lat0 + self.dlat and
                ll[:, 0].max() >= self.lon0 - self.dlon and ll[:, 0].min() <= self.lon0 + self.dlon)


def line_items(feats, fr, eps_m, counter):
    """Project, clip, simplify (metres), quantise, dedupe; returns [(key, qpts, km_len_in_square)]; counts collapses."""
    out = []
    for f in feats:
        if not fr.bbox_hit(f["ll"]):
            continue
        for ri, run in enumerate(clip_line(fr.es(f["ll"]), fr.R)):
            km = sum(math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(run, run[1:]))
            simp = dp([(x * 1000, y * 1000) for x, y in run], eps_m)
            q = quant([(x / 1000, y / 1000) for x, y in simp], fr.R_mm)
            if len(q) < 2:
                counter[0] += 1
                continue
            out.append((tuple(f["key"]) + (ri,), q, km, f))
    out.sort(key=lambda t: t[0])
    return out


def disc_len_m(pts, R):
    """RD8: the raw clipped length inside the disc of radius R, in metres."""
    tot = 0.0
    for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
        dx, dy = x1 - x0, y1 - y0
        a = dx * dx + dy * dy
        if a == 0:
            continue
        b = 2 * (x0 * dx + y0 * dy)
        c = x0 * x0 + y0 * y0 - R * R
        disc = b * b - 4 * a * c
        if disc <= 0:
            continue
        sq = math.sqrt(disc)
        t0, t1 = max(0.0, (-b - sq) / (2 * a)), min(1.0, (-b + sq) / (2 * a))
        if t1 > t0:
            tot += (t1 - t0) * math.sqrt(a) * 1000
    return tot


# ---------------------------------------------------------------------------------------------- roads: chaining (Q18)
def chain(roads):
    """Links meeting end to end, with the same road number and class, at a node where exactly two such links meet, become
    one polyline. Deterministic: seeds in source order, extended forward then backward. Unnumbered links are not chained."""
    by_node = {}
    for i, r in enumerate(roads):
        if r["num"] and r["a"] and r["b"]:
            for n in (r["a"], r["b"]):
                by_node.setdefault((r["num"], r["cls"], n), []).append(i)
    used = [False] * len(roads)
    out = []
    for i in sorted(range(len(roads)), key=lambda k: roads[k]["key"]):
        if used[i]:
            continue
        used[i] = True
        r = roads[i]
        if not (r["num"] and r["a"] and r["b"]):
            out.append(r)
            continue
        seq = [(i, False)]
        for direction in ("fwd", "back"):
            node = r["b"] if direction == "fwd" else r["a"]
            while True:
                ids = by_node.get((r["num"], r["cls"], node), [])
                if len(ids) != 2:
                    break
                nxt = [k for k in ids if not used[k]]
                if len(nxt) != 1:
                    break
                k = nxt[0]
                used[k] = True
                rk = roads[k]
                if direction == "fwd":
                    rev = rk["b"] == node
                    seq.append((k, rev))
                    node = rk["a"] if rev else rk["b"]
                else:
                    rev = rk["a"] == node
                    seq.insert(0, (k, rev))
                    node = rk["b"] if rev else rk["a"]
        coords = []
        for k, rev in seq:
            c = roads[k]["ll"][::-1] if rev else roads[k]["ll"]
            coords.extend(map(tuple, c if not coords else c[1:]))
        out.append({**r, "ll": np.asarray(coords), "members": len(seq)})
    return out


# ---------------------------------------------------------------------------------------------- terrain z9 (RD10, RD17, RD21)
_tiles, _sources = {}, {}


def tile(x, y):
    if (x, y) in _tiles:
        return _tiles[(x, y)]
    d = os.path.join(SRC, "terrain", "9", str(x))
    p = os.path.join(d, f"{y}.png")
    meta = p + ".json"
    if not os.path.exists(p):
        os.makedirs(d, exist_ok=True)
        url = f"https://s3.amazonaws.com/elevation-tiles-prod/terrarium/9/{x}/{y}.png"
        with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
            b = r.read()
            hdr = {"url": url, "sha256": hashlib.sha256(b).hexdigest(), "last_modified": r.headers.get("Last-Modified"),
                   "sources": r.headers.get("x-amz-meta-x-imagery-sources")}
        open(p, "wb").write(b)
        json.dump(hdr, open(meta, "w"))
    a = np.asarray(Image.open(p).convert("RGB"), dtype=np.float64)
    _tiles[(x, y)] = a[:, :, 0] * 256 + a[:, :, 1] + a[:, :, 2] / 256 - 32768
    _sources[(x, y)] = json.load(open(meta))
    return _tiles[(x, y)]


def pix(lon, lat):
    phi = lat * math.pi / 180
    X = math.floor((lon + 180) / 360 * 2 ** 17)
    Y = math.floor((1 - math.log(math.tan(math.pi / 4 + phi / 2)) / math.pi) / 2 * 2 ** 17)
    return X, Y


def elev_px(X, Y):
    return tile(X // 256, Y // 256)[Y % 256, X % 256]


def elev(lon, lat):
    return math.floor(elev_px(*pix(lon, lat)) + 0.5)


def max_within_1km(lon, lat, fr):
    X, Y = pix(lon, lat)
    e0, s0 = fr.es([(lon, lat)])[0]
    best = -1e9
    for dy in range(-8, 9):
        for dx in range(-8, 9):
            px, py = X + dx, Y + dy
            plon = (px + 0.5) / 2 ** 17 * 360 - 180
            n = math.pi - 2 * math.pi * (py + 0.5) / 2 ** 17
            plat = math.degrees(math.atan(math.sinh(n)))
            e, s = fr.es([(plon, plat)])[0]
            if math.hypot(e - e0, s - s0) <= 1.0:
                best = max(best, elev_px(px, py))
    return math.floor(best + 0.5)


# ---------------------------------------------------------------------------------------------- one package
def build(gb, ni, centre, lat, lon, r_km, eps_req, variant, ni_in, nodrop=False):
    fr = Frame(lat, lon, r_km)
    collapsed = {k: [0] for k in ("water", "rivers", "roads", "peaks", "places")}
    ni_on = ni_in and ni is not None

    def roads_for(eps):
        cls2_primary_only = variant in ("ii", "iv")
        sel = [r for r in gb["roads"] if r["cls"] == 1 or (not cls2_primary_only or r["primary"])]
        sel = [r for r in sel if fr.bbox_hit(r["ll"])]
        if variant in ("iii", "iv") and "--nochain" not in FLAGS:
            sel = chain(sel)
        if ni_on:
            sel = sel + [r for r in ni["roads"] if fr.bbox_hit(r["ll"])]
        return line_items(sel, fr, eps, collapsed["roads"])

    def water_for(eps):
        feats = gb["water"] + (ni["water"] if ni_on else [])
        out = []
        for f in feats:
            if not fr.bbox_hit(f["ll"]):
                continue
            ring = clip_ring(fr.es(f["ll"]), fr.R)
            if len(ring) < 3:
                continue
            path = [(x * 1000, y * 1000) for x, y in ring]
            simp = dp(path + [path[0]], eps)[:-1]
            q = quant([(x / 1000, y / 1000) for x, y in simp], fr.R_mm)
            if len(q) > 1 and q[-1] == q[0]:
                q = q[:-1]
            if len(q) < 3:
                collapsed["water"][0] += 1
                continue
            out.append((tuple(f["key"]), q))
        out.sort(key=lambda t: t[0])
        return out

    def rivers_for(eps):
        groups = {}
        for f in gb["rivers"]:
            if fr.bbox_hit(f["ll"]):
                groups.setdefault(f["name"], []).append(f)
        lens = {n: sum(disc_len_m(fr.es(f["ll"]), fr.R) for f in fs) for n, fs in groups.items()}
        kept = sorted([n for n in lens if lens[n] > 0], key=lambda n: (-lens[n], n.encode("utf-8")))[:5]
        names = kept
        items = []
        for idx, n in enumerate(names):
            for key, q, km, f in line_items(groups[n], fr, eps, collapsed["rivers"]):
                items.append((key, idx, q))
        items.sort(key=lambda t: t[0])
        return names, items

    def peaks_for():
        cand = []
        for f in gb["hills"]:
            e, s = fr.es([(f["lon"], f["lat"])])[0]
            if math.hypot(e, s) <= fr.R:
                cand.append((elev(f["lon"], f["lat"]), f["id"], f, e, s))
        cand.sort(key=lambda t: (-t[0], t[1].encode("utf-8")))
        kept = []
        for el, fid, f, e, s in cand:
            if any(math.hypot(e - k[3], s - k[4]) <= 3.0 for k in kept):
                continue
            kept.append((el, fid, f, e, s))
            if len(kept) == 16:
                break
        out, q16 = [], []
        for el, fid, f, e, s in kept:
            qx, qy = quant([(e, s)], fr.R_mm)[0]
            out.append((qx, qy, int(el), f["name"]))
            q16.append(max_within_1km(f["lon"], f["lat"], fr) - el)
        return out, q16

    def places_for():
        out = []
        feats = gb["places"] + (ni["places"] if ni_on else [])
        for f in sorted(feats, key=lambda f: f["key"]):
            e, s = fr.es([(f["lon"], f["lat"])])[0]
            if math.hypot(e, s) <= fr.R:
                qx, qy = quant([(e, s)], fr.R_mm)[0]
                out.append((qx, qy, int(f["pop"]), f["name"]))
        return out

    peaks, q16 = peaks_for()
    places = places_for()

    def provenance(eps_req, dropped, coll, files):
        lines = ["format: BMPK 1", "builder: m1_sizes 000000000000", "store: 000000000000"]
        for layer in ("water", "rivers", "roads", "peaks", "places", "population", "terrain"):
            fl = files.get(layer, [])
            if not fl:
                continue
            pre = os.path.commonprefix(fl)
            pre = pre[: pre.rfind("/") + 1]
            sha = hashlib.sha256("".join(f"{u} {'0' * 64}\n" for u in sorted(fl)).encode()).hexdigest()
            lines += [f"source.{layer}: {pre}", f"vintage.{layer}: 2026", f"files.{layer}: {len(fl)}", f"sha256.{layer}: {sha}"]
        lines += ["credit.terrain: Produced using Copernicus data and information funded by the European Union - EU-DEM layers",
                  "credit.terrain: SRTM data courtesy of the U.S. Geological Survey",
                  "credit.terrain: ETOPO1: DOC/NOAA/NESDIS/NCEI > National Centers for Environmental Information, NESDIS, NOAA, U.S. Department of Commerce",
                  "processing.terrain: Terrarium z9 elevations; summit elevations taken from the z9 pixel containing the summit, rounded to whole metres"]
        lines.append(f"epsilon.m: {eps_req:.3f}")   # RD-P4: the requested mm / 1,000, three decimals
        for layer in ("water", "rivers", "roads", "peaks", "places"):
            lines += [f"dropped.{layer}: {dropped.get(layer, 0)}", f"collapsed.{layer}: {coll.get(layer, 0)}"]
        return lines

    def encode(eps_used, water, rivers, roads, dropped):
        names, ritems = rivers
        files = {
            "water": sorted({w[0][0] for w in water}), "rivers": sorted({it[0][0] for it in ritems}),
            "roads": sorted({r[0][0] for r in roads}), "peaks": ["os-names/Data"] if peaks else [],
            "places": ["os-names/Data"] if places else [], "population": ["ons", "ons-bua", "nrs", "nisra"],
            "terrain": [f"terrain/9/{x}/{y}.png" for (x, y) in sorted(_tiles)],
        }
        coll = {k: v[0] for k, v in collapsed.items()}
        secs = [
            (bmpk.WATER, len(water), bmpk.water_body([w[1] for w in water])),
            (bmpk.RIVERS, len(ritems), bmpk.rivers_body(names, [(it[1], it[2]) for it in ritems])),
            (bmpk.ROADS, len(roads), bmpk.roads_body([(r[3]["cls"], r[1]) for r in roads])),
            (bmpk.PEAKS, len(peaks), bmpk.peaks_body(peaks)),
            (bmpk.PLACES, len(places), bmpk.places_body(places)),
        ]
        prov = provenance(eps_req, dropped, coll, files)
        secs.append((bmpk.PROVENANCE, len(prov), bmpk.provenance_body(prov)))
        pkg = bmpk.encode(secs, math.floor(fr.lat0 * 1e6 + 0.5), math.floor(fr.lon0 * 1e6 + 0.5), fr.R_mm,
                          math.floor(eps_used * 1000 + 0.5))
        return pkg, {s[0]: len(s[2]) for s in secs}

    eps = 0.0 if "--eps0" in FLAGS else eps_req
    water, rivers, roads = water_for(eps), rivers_for(eps), roads_for(eps)
    pkg, secl = encode(eps, water, rivers, roads, {})
    first = len(pkg)
    cls2_km_all = sum(r[2] for r in roads if r[3]["cls"] == 2)
    outcome, dropped, doubled = ("F" if first <= F_LIMIT else "C") if first <= CAP else None, {}, False
    if outcome is None:
        if nodrop:
            outcome = "X"
        else:
            doubled = True
            eps2 = eps * 2
            for k in collapsed:
                collapsed[k][0] = 0
            water, rivers, roads = water_for(eps2), rivers_for(eps2), roads_for(eps2)
            pkg, secl = encode(eps2, water, rivers, roads, {})
            if len(pkg) > CAP:
                wsz = sorted(range(len(water)), key=lambda i: (abs(sum(x0 * y1 - x1 * y0 for (x0, y0), (x1, y1) in
                                                                     zip(water[i][1], water[i][1][1:] + water[i][1][:1]))), -i))
                rsz = sorted([i for i in range(len(roads)) if roads[i][3]["cls"] == 2],
                             key=lambda i: (sum(math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(roads[i][1], roads[i][1][1:])), -i))
                drop_w, drop_r = set(), set()
                order = [("water", i) for i in wsz] + [("roads", i) for i in rsz]
                lo, hi = 0, len(order)
                # the smallest prefix of the drop order that fits: bisect on the exact encoded size
                def fits(n):
                    dw = {i for l, i in order[:n] if l == "water"}
                    dr = {i for l, i in order[:n] if l == "roads"}
                    p, s = encode(eps2, [w for i, w in enumerate(water) if i not in dw], rivers,
                                  [r for i, r in enumerate(roads) if i not in dr],
                                  {"water": len(dw), "roads": len(dr)})
                    return len(p) <= CAP, p, s, dw, dr
                ok_all = fits(hi)
                if not ok_all[0]:
                    outcome = "X"
                else:
                    while lo < hi:
                        mid = (lo + hi) // 2
                        if fits(mid)[0]:
                            hi = mid
                        else:
                            lo = mid + 1
                    ok, pkg, secl, drop_w, drop_r = fits(lo)
                    roads = [r for i, r in enumerate(roads) if i not in drop_r]
                    water = [w for i, w in enumerate(water) if i not in drop_w]
                    dropped = {"water": len(drop_w), "roads": len(drop_r)}
                    outcome = "D"
            else:
                outcome = "D"
    st, info = bmpk_validate.validate(pkg)
    cls2_km_kept = sum(r[2] for r in roads if r[3]["cls"] == 2)
    return {
        "centre": centre, "r_km": r_km, "eps_m": eps_req, "variant": variant, "ni_in": ni_on, "outcome": outcome,
        "bytes_first": first, "bytes_final": len(pkg), "validate": st, "eps_doubled": doubled, "dropped": dropped,
        "collapsed": {k: v[0] for k, v in collapsed.items()}, "section_bytes": {str(k): v for k, v in secl.items()},
        "counts": {"water": len(water), "rivers": len(rivers[1]), "river_groups": len(rivers[0]), "roads": len(roads),
                   "peaks": len(peaks), "places": len(places)},
        "class2_km_kept_share": round(cls2_km_kept / cls2_km_all, 4) if cls2_km_all else None,
        "q16_diffs_m": q16,
    }


# ---------------------------------------------------------------------------------------------- controls (A2.4)
def controls():
    out = {}
    r = subprocess.run([sys.executable, "-I", os.path.join(REPO, "tools", "weatherscope", "basemap", "synth.py"),
                        os.path.join(REPO, "test", "fixtures", "wx", "basemap"), "--check"], capture_output=True, text=True)
    out["synth_check"] = {"exit": r.returncode, "tail": (r.stdout + r.stderr).strip().splitlines()[-2:]}
    eps = 100.0
    line = [(0.0, 0.0), (500.0, eps / 2), (1000.0, 0.0)]
    out["dp_bump_half_eps"] = len(dp(line, eps)) == 2
    out["dp_bump_double_eps_kept"] = len(dp([(0.0, 0.0), (500.0, 2 * eps), (1000.0, 0.0)], eps)) == 3
    ring = clip_ring([(-2.0, -0.5), (0.5, -0.5), (0.5, 0.5), (-2.0, 0.5)], 1.0)
    out["ring_clipped_closed"] = len(ring) >= 3 and all(-1.0 <= x <= 1.0 and -1.0 <= y <= 1.0 for x, y in ring)
    runs = clip_line([(-2.0, 0.0), (2.0, 0.0), (2.0, 0.5), (-2.0, 0.5)], 1.0)
    out["line_reentry_two_runs"] = len(runs) == 2
    return out


def main():
    apts = {}
    with open(APTS, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            c = row.get("iata_code") or ""
            if c in GB_CENTRES + NI_CENTRES and row["type"] in ("large_airport", "medium_airport", "small_airport"):
                apts[c] = (float(row["latitude_deg"]), float(row["longitude_deg"]))
    t0 = time.time()
    res = {"flags": sorted(FLAGS), "controls": controls()}
    gb, ni = load_gb(), load_ni()
    res["pop_stats"] = {"gb": gb["pop_stats"], "ni": ni["pop_stats"]}
    runs = []
    for c in GB_CENTRES + NI_CENTRES:
        lat, lon = apts[c]
        for r_km, eps in RADII:
            for v in VARIANTS:
                runs.append(build(gb, ni, c, lat, lon, r_km, eps, v, True, nodrop="--nodrop" in FLAGS))
                print(json.dumps({k: runs[-1][k] for k in ("centre", "r_km", "variant", "outcome", "bytes_first", "bytes_final", "validate")}), flush=True)
        if c in ("BFS", "LDY", "GLA"):
            for r_km, eps in RADII[:2]:
                runs.append(build(gb, ni, c, lat, lon, r_km, eps, "iv", False, nodrop="--nodrop" in FLAGS))
    res["runs"] = runs
    res["terrain_tiles"] = len(_tiles)
    res["terrain_sources"] = sorted({s for v in _sources.values() for s in (v.get("sources") or "").split(",") if s})[:40]
    res["seconds"] = round(time.time() - t0)
    json.dump(res, open(OUT, "w"), indent=1)
    print(json.dumps({"controls": res["controls"], "seconds": res["seconds"], "terrain_tiles": res["terrain_tiles"]}))


main()
